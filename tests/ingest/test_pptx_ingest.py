import io
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.presentation import Presentation as PresentationType
from pptx.shapes.autoshape import Shape
from pptx.slide import Slide
from pptx.util import Inches

from alfred.ingest import BlockKind, IngestError, Limits, SourceFormat, Unit, ingest
from alfred.ingest.errors import FileTooLargeError, TooManyPagesError
from alfred.ingest.pptx_parser import HIDDEN_ISSUE, IMAGE_ONLY_ISSUE, equation_issue

TITLE_AND_CONTENT = 1  # index of the "Title and Content" layout in the default template
BLANK = 6


def save(presentation: PresentationType, tmp_path: Path) -> Path:
    path = tmp_path / "deck.pptx"
    presentation.save(str(path))
    return path


def add_textbox(slide: Slide, text: str, *, top: float) -> None:
    box = slide.shapes.add_textbox(Inches(1), Inches(top), Inches(6), Inches(1))
    box.text_frame.text = text


def add_picture(slide: Slide) -> None:
    image = io.BytesIO()
    Image.new("RGB", (8, 8)).save(image, "PNG")
    image.seek(0)
    slide.shapes.add_picture(image, Inches(1), Inches(1))


def test_extracts_title_nested_bullets_and_notes(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[TITLE_AND_CONTENT])
    assert slide.shapes.title is not None
    slide.shapes.title.text = "Self-attention"
    body_shape = slide.placeholders[1]
    assert isinstance(body_shape, Shape)
    body = body_shape.text_frame
    body.text = "Queries, keys and values"
    sub_bullet = body.add_paragraph()
    sub_bullet.text = "All three are projections of the input"
    sub_bullet.level = 1
    notes = slide.notes_slide.notes_text_frame
    assert notes is not None
    notes.text = "Mention the scaling by sqrt(d_k)."

    document = ingest(save(presentation, tmp_path))

    assert document.format is SourceFormat.PPTX
    assert document.unit is Unit.SLIDE
    assert document.parser.startswith("python-pptx-")
    blocks = document.page(1).blocks
    assert [b.kind for b in blocks] == [BlockKind.TITLE, BlockKind.TEXT, BlockKind.NOTES]
    assert blocks[0].text == "Self-attention"
    assert blocks[1].text == "Queries, keys and values\n  All three are projections of the input"
    assert blocks[2].text == "Mention the scaling by sqrt(d_k)."


def test_tables_become_rows_of_cells(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[BLANK])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(1)).table
    for (row, col), text in {
        (0, 0): "Model",
        (0, 1): "Params",
        (1, 0): "BERT",
        (1, 1): "110M",
    }.items():
        table.cell(row, col).text = text

    page = ingest(save(presentation, tmp_path)).page(1)

    assert page.blocks[0].kind is BlockKind.TABLE
    assert page.blocks[0].text == "Model | Params\nBERT | 110M"


def test_text_follows_reading_order_not_creation_order(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[BLANK])
    add_textbox(slide, "bottom", top=5)
    add_textbox(slide, "top", top=1)

    page = ingest(save(presentation, tmp_path)).page(1)

    assert [b.text for b in page.blocks] == ["top", "bottom"]


def test_reads_text_inside_groups(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[BLANK])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "grouped"

    page = ingest(save(presentation, tmp_path)).page(1)

    assert page.text == "grouped"


def test_flags_slides_that_are_mostly_images(tmp_path: Path) -> None:
    presentation = Presentation()
    image_only = presentation.slides.add_slide(presentation.slide_layouts[BLANK])
    add_picture(image_only)
    add_textbox(image_only, "Photo: CC BY", top=6)
    with_text = presentation.slides.add_slide(presentation.slide_layouts[BLANK])
    add_picture(with_text)
    add_textbox(with_text, "A diagram of the encoder-decoder architecture.", top=6)

    document = ingest(save(presentation, tmp_path))

    assert IMAGE_ONLY_ISSUE in document.page(1).issues
    assert IMAGE_ONLY_ISSUE not in document.page(2).issues


def test_hidden_slides_are_kept_and_flagged(tmp_path: Path) -> None:
    presentation = Presentation()
    for text in ("first", "hidden", "third"):
        add_textbox(presentation.slides.add_slide(presentation.slide_layouts[BLANK]), text, top=1)
    presentation.slides[1].element.set("show", "0")

    document = ingest(save(presentation, tmp_path))

    assert [p.text for p in document.pages] == ["first", "hidden", "third"]
    assert document.page(2).issues == (HIDDEN_ISSUE,)


def test_flags_equations(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[BLANK])
    add_textbox(slide, "The cosine similarity is defined as follows.", top=1)
    equation = parse_xml(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        "<m:r><m:t>cos</m:t></m:r></m:oMath>"
    )
    textbox = slide.shapes[0]
    assert isinstance(textbox, Shape)
    textbox.text_frame.paragraphs[0]._p.append(equation)

    page = ingest(save(presentation, tmp_path)).page(1)

    assert equation_issue(1) in page.issues


def test_rejects_decks_over_the_slide_limit(tmp_path: Path) -> None:
    presentation = Presentation()
    for _ in range(3):
        presentation.slides.add_slide(presentation.slide_layouts[BLANK])

    with pytest.raises(TooManyPagesError):
        ingest(save(presentation, tmp_path), Limits(max_pages=2))


def test_rejects_archives_that_unpack_too_large(tmp_path: Path) -> None:
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[BLANK])

    with pytest.raises(FileTooLargeError, match="unpacks to"):
        ingest(save(presentation, tmp_path), Limits(max_unpacked_bytes=1000))


def test_damaged_pptx_raises_ingest_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.pptx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("ppt/presentation.xml", "this is not xml")

    with pytest.raises(IngestError, match="Could not open"):
        ingest(path)
