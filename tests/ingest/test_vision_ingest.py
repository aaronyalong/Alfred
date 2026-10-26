import base64
import io
import json
from pathlib import Path
from types import SimpleNamespace
from typing import Any, cast

import anthropic
import pytest
from fpdf import FPDF
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from alfred.ingest import BlockKind, Engine, ingest
from alfred.ingest.issues import visual_issue
from alfred.ingest.pptx_parser import HIDDEN_ISSUE
from alfred.ingest.vision import (
    DECLINED_ISSUE,
    FAILED_ISSUE,
    MAX_IMAGE_SIDE,
    VISION_ISSUE,
    VisionCache,
    VisionReader,
    VisionUsage,
)

TEXT = "This page has enough text to count as a real page of content."


class FakeMessages:
    """Stands in for client.beta.messages: records each request, returns a canned reply."""

    def __init__(self, reply: str, stop_reason: str, model: str) -> None:
        self.requests: list[dict[str, Any]] = []
        self._response = SimpleNamespace(
            model=model,
            stop_reason=stop_reason,
            usage=SimpleNamespace(input_tokens=1_000, output_tokens=200),
            content=[SimpleNamespace(type="text", text=reply)],
        )

    def create(self, **request: Any) -> SimpleNamespace:
        self.requests.append(request)
        return self._response


def fake_reader(
    blocks: list[dict[str, str]] | None = None,
    *,
    reply: str | None = None,
    stop_reason: str = "end_turn",
    model: str = "claude-opus-5-5",
    cache: VisionCache | None = None,
) -> tuple[VisionReader, FakeMessages]:
    reply = json.dumps({"blocks": blocks or []}) if reply is None else reply
    messages = FakeMessages(reply, stop_reason, model)
    client = SimpleNamespace(beta=SimpleNamespace(messages=messages))
    return VisionReader(cast(anthropic.Anthropic, client), model=model, cache=cache), messages


def sent_images(request: dict[str, Any]) -> list[Image.Image]:
    content = request["messages"][0]["content"]
    return [
        Image.open(io.BytesIO(base64.b64decode(block["source"]["data"])))
        for block in content
        if block["type"] == "image"
    ]


def sent_prompt(request: dict[str, Any]) -> str:
    return request["messages"][0]["content"][-1]["text"]


def pdf_with_image_on_page_2(tmp_path: Path) -> Path:
    pdf = FPDF()
    pdf.set_font("helvetica", size=12)
    for draw_image in (False, True):
        pdf.add_page()
        pdf.multi_cell(0, 8, TEXT)
        if draw_image:
            pdf.image(Image.new("RGB", (8, 8)), x=10, y=40, w=150, h=100)
    path = tmp_path / "lecture.pdf"
    pdf.output(str(path))
    return path


def test_only_flagged_pdf_pages_are_sent_and_get_extra_blocks(tmp_path: Path) -> None:
    reader, messages = fake_reader(
        [
            {"kind": "formula", "text": r"\cos\theta = \frac{a \cdot b}{|a||b|}"},
            {"kind": "diagram", "text": "Two vectors a and b with angle theta. Labels: a, b"},
            {"kind": "text", "text": "  "},  # empty blocks are dropped
        ]
    )

    document = ingest(pdf_with_image_on_page_2(tmp_path), engine=Engine.VISION, vision=reader)

    assert len(messages.requests) == 1
    assert TEXT in sent_prompt(messages.requests[0])
    [image] = sent_images(messages.requests[0])
    assert max(image.size) == MAX_IMAGE_SIDE
    assert document.page(1).issues == ()
    page = document.page(2)
    assert [b.kind for b in page.blocks] == [BlockKind.TEXT, BlockKind.FORMULA, BlockKind.DIAGRAM]
    assert page.issues == (VISION_ISSUE,)
    assert document.parser.endswith("+vision:claude-opus-5-5")


def test_slides_send_their_pictures_and_keep_issues_pictures_cannot_resolve(
    tmp_path: Path,
) -> None:
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    hidden = presentation.slides.add_slide(blank)
    hidden.element.set("show", "0")
    diagram = presentation.slides.add_slide(blank)
    picture = io.BytesIO()
    Image.new("RGB", (40, 30)).save(picture, "PNG")
    diagram.shapes.add_picture(picture, Inches(1), Inches(1), Inches(5), Inches(4))
    path = tmp_path / "deck.pptx"
    presentation.save(str(path))
    reader, messages = fake_reader([{"kind": "diagram", "text": "A VPC with two subnets."}])

    document = ingest(path, engine=Engine.VISION, vision=reader)

    assert len(messages.requests) == 1  # the hidden slide has nothing missing
    assert [img.size for img in sent_images(messages.requests[0])] == [(40, 30)]
    assert "pictures on this slide" in sent_prompt(messages.requests[0])
    assert document.page(1).issues == (HIDDEN_ISSUE,)
    assert document.page(2).issues == (VISION_ISSUE,)
    assert document.page(2).blocks[-1].kind is BlockKind.DIAGRAM


def test_a_declined_page_keeps_its_baseline_content_and_issues(tmp_path: Path) -> None:
    reader, _ = fake_reader(reply="I can't help with that.", stop_reason="refusal")

    page = ingest(pdf_with_image_on_page_2(tmp_path), engine=Engine.VISION, vision=reader).page(2)

    assert page.text == TEXT
    assert page.issues == (visual_issue(1), DECLINED_ISSUE)
    assert reader.usage.calls == 1  # declined calls still count


def test_a_reading_cut_off_at_max_tokens_is_flagged_not_raised(tmp_path: Path) -> None:
    reader, _ = fake_reader(reply='{"blocks": [{"kind": "formula", "te', stop_reason="max_tokens")

    page = ingest(pdf_with_image_on_page_2(tmp_path), engine=Engine.VISION, vision=reader).page(2)

    assert page.text == TEXT
    assert page.issues == (visual_issue(1), FAILED_ISSUE)
    assert reader.usage.output_tokens == 200


def test_documents_without_flagged_pages_make_no_calls(make_pdf: Any) -> None:
    reader, messages = fake_reader()

    ingest(make_pdf([TEXT]), engine=Engine.VISION, vision=reader)

    assert messages.requests == []
    assert reader.usage.calls == 0


def test_usage_adds_up_tokens_and_cost(tmp_path: Path) -> None:
    reader, _ = fake_reader()

    ingest(pdf_with_image_on_page_2(tmp_path), engine=Engine.VISION, vision=reader)

    usage = reader.usage
    assert (usage.calls, usage.input_tokens, usage.output_tokens) == (1, 1_000, 200)
    assert usage.cost_usd == pytest.approx((1_000 * 4.00 + 200 * 20.00) / 1e6)
    assert usage.served_by == {"claude-opus-5-5"}


def test_cost_is_unknown_once_a_call_ran_on_an_unpriced_model() -> None:
    usage = VisionUsage()

    usage.add("claude-opus-5-5", 1_000, 200)
    usage.add("some-fallback-model", 1_000, 200)

    assert usage.cost_usd is None
    assert usage.calls == 2


FORMULA = [{"kind": "formula", "text": r"\sqrt{d_k}"}]


def test_a_cached_reading_is_reused_without_a_call(tmp_path: Path) -> None:
    pdf, cache = pdf_with_image_on_page_2(tmp_path), VisionCache(tmp_path / "cache")
    first, _ = fake_reader(FORMULA, cache=cache)
    second, messages = fake_reader(cache=cache)  # would return no blocks if it were called

    expected = ingest(pdf, engine=Engine.VISION, vision=first)
    document = ingest(pdf, engine=Engine.VISION, vision=second)

    assert messages.requests == []
    assert (second.usage.calls, second.usage.cached, second.usage.cost_usd) == (0, 1, 0.0)
    assert document.pages == expected.pages


def test_failed_readings_are_not_cached(tmp_path: Path) -> None:
    pdf, cache = pdf_with_image_on_page_2(tmp_path), VisionCache(tmp_path / "cache")
    declined, _ = fake_reader(reply="No.", stop_reason="refusal", cache=cache)
    retry, messages = fake_reader(FORMULA, cache=cache)

    ingest(pdf, engine=Engine.VISION, vision=declined)
    page = ingest(pdf, engine=Engine.VISION, vision=retry).page(2)

    assert len(messages.requests) == 1
    assert page.issues == (VISION_ISSUE,)


def test_a_different_model_does_not_reuse_cached_readings(tmp_path: Path) -> None:
    pdf, cache = pdf_with_image_on_page_2(tmp_path), VisionCache(tmp_path / "cache")
    opus, _ = fake_reader(FORMULA, cache=cache)
    sonnet, messages = fake_reader(FORMULA, model="claude-sonnet-5-5", cache=cache)

    ingest(pdf, engine=Engine.VISION, vision=opus)
    ingest(pdf, engine=Engine.VISION, vision=sonnet)

    assert len(messages.requests) == 1


def test_an_unreadable_cache_file_counts_as_a_miss(tmp_path: Path) -> None:
    pdf, cache_dir = pdf_with_image_on_page_2(tmp_path), tmp_path / "cache"
    first, _ = fake_reader(FORMULA, cache=VisionCache(cache_dir))
    ingest(pdf, engine=Engine.VISION, vision=first)
    for saved in cache_dir.iterdir():
        saved.write_text("{not json", encoding="utf-8")
    second, messages = fake_reader(FORMULA, cache=VisionCache(cache_dir))

    ingest(pdf, engine=Engine.VISION, vision=second)

    assert len(messages.requests) == 1
