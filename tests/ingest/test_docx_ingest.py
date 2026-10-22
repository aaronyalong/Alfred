import io
import zipfile
from pathlib import Path

import pytest
from docx import Document as new_docx
from docx.document import Document as DocxDocument
from docx.oxml import parse_xml
from PIL import Image

from alfred.ingest import BlockKind, IngestError, Limits, SourceFormat, Unit, ingest
from alfred.ingest.docx_parser import MAX_SECTION_CHARS
from alfred.ingest.errors import TooManyPagesError
from alfred.ingest.issues import equation_issue, visual_issue


def save(document: DocxDocument, tmp_path: Path) -> Path:
    path = tmp_path / "notes.docx"
    document.save(str(path))
    return path


def test_headings_start_new_sections(tmp_path: Path) -> None:
    document = new_docx()
    document.add_heading("Virtualization", level=1)
    document.add_paragraph("A hypervisor runs multiple virtual machines on one host.")
    document.add_heading("Containers", level=1)
    document.add_paragraph("Containers share the host kernel.")

    parsed = ingest(save(document, tmp_path))

    assert parsed.format is SourceFormat.DOCX
    assert parsed.unit is Unit.SECTION
    assert parsed.parser.startswith("python-docx-")
    assert len(parsed.pages) == 2
    first = parsed.page(1).blocks
    assert [b.kind for b in first] == [BlockKind.TITLE, BlockKind.TEXT]
    assert first[0].text == "Virtualization"
    assert parsed.page(2).blocks[0].text == "Containers"


def test_without_headings_long_documents_split_by_size(tmp_path: Path) -> None:
    document = new_docx()
    paragraphs = [f"Paragraph {i}: " + "x" * 985 for i in range(10)]  # ~1,000 chars each
    for text in paragraphs:
        document.add_paragraph(text)

    parsed = ingest(save(document, tmp_path))

    assert len(parsed.pages) == 4  # 3 + 3 + 3 + 1 paragraphs
    assert all(sum(len(b.text) for b in p.blocks) <= MAX_SECTION_CHARS for p in parsed.pages)
    extracted = [b.text for p in parsed.pages for b in p.blocks]
    assert extracted == paragraphs  # nothing lost, nothing reordered, nothing cut


def test_a_heading_like_line_moves_to_the_section_it_introduces(tmp_path: Path) -> None:
    document = new_docx()
    document.add_paragraph("x" * 2000 + ".")
    document.add_paragraph("Part III: Application Problems (30%)")  # plain text, no style
    document.add_paragraph("y" * 2000 + ".")

    parsed = ingest(save(document, tmp_path))

    assert len(parsed.pages) == 2
    assert parsed.page(2).blocks[0].text == "Part III: Application Problems (30%)"


def test_a_paragraph_longer_than_a_section_is_never_cut(tmp_path: Path) -> None:
    document = new_docx()
    long_paragraph = "word " * 1000  # 5,000 chars
    document.add_paragraph(long_paragraph)

    parsed = ingest(save(document, tmp_path))

    assert len(parsed.pages) == 1
    assert parsed.page(1).text == long_paragraph.strip()


def test_list_items_are_marked_and_indented(tmp_path: Path) -> None:
    document = new_docx()
    document.add_paragraph("Service models", style="List Bullet")
    document.add_paragraph("Infrastructure as a Service", style="List Bullet 2")
    numbered = document.add_paragraph("Platform as a Service")
    numbered._p.get_or_add_pPr().append(
        parse_xml(
            '<w:numPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:ilvl w:val="1"/><w:numId w:val="1"/></w:numPr>'
        )
    )

    texts = [b.text for b in ingest(save(document, tmp_path)).page(1).blocks]

    assert texts == [
        "- Service models",
        "  - Infrastructure as a Service",
        "  - Platform as a Service",
    ]


def test_tables_become_rows_and_merged_cells_appear_once(tmp_path: Path) -> None:
    document = new_docx()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Cloud service models"
    for col, text in enumerate(["IaaS", "PaaS", "SaaS"]):
        table.cell(1, col).text = text

    block = ingest(save(document, tmp_path)).page(1).blocks[0]

    assert block.kind is BlockKind.TABLE
    assert block.text == "Cloud service models\nIaaS | PaaS | SaaS"


def test_flags_images_and_equations(tmp_path: Path) -> None:
    document = new_docx()
    document.add_paragraph("The diagram below shows a three-tier architecture.")
    image = io.BytesIO()
    Image.new("RGB", (8, 8)).save(image, "PNG")
    image.seek(0)
    document.add_picture(image)
    with_equation = document.add_paragraph("Availability is computed as:")
    with_equation._p.append(
        parse_xml(
            '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            "<m:r><m:t>A</m:t></m:r></m:oMath>"
        )
    )

    page = ingest(save(document, tmp_path)).page(1)

    assert page.issues == (visual_issue(1), equation_issue(1))


def test_rejects_documents_with_too_many_sections(tmp_path: Path) -> None:
    document = new_docx()
    for title in ("One", "Two", "Three"):
        document.add_heading(title, level=1)

    with pytest.raises(TooManyPagesError):
        ingest(save(document, tmp_path), Limits(max_pages=2))


def test_damaged_docx_raises_ingest_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", "this is not xml")

    with pytest.raises(IngestError, match="Could not open"):
        ingest(path)
