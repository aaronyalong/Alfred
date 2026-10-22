"""Baseline PPTX parser: titles, text, tables and speaker notes per slide, via python-pptx.

Unlike a PDF export, the PPTX file still knows which text is the title, how bullets are
nested, and what the speaker notes say. Not extracted yet: equations, charts, SmartArt and
text inside images. Slides that contain them are flagged, so the gap is visible, not silent.
"""

from collections.abc import Iterable, Iterator
from importlib.metadata import version
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture
from pptx.slide import Slide

from alfred.ingest.errors import IngestError, TooManyPagesError
from alfred.ingest.models import Block, BlockKind, Page
from alfred.ingest.text import clean_text

PARSER_NAME = f"python-pptx-{version('python-pptx')}"

HIDDEN_ISSUE = "hidden slide"

# Pictures and charts covering at least this share of the slide are treated as content
# (diagrams, plots), not decoration. Measured on a real cloud deck: every diagram covered
# 27-60% of its slide, while icons and logos stay well below 10%.
MIN_VISUAL_SHARE = 0.10

_TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
_MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def visual_issue(count: int) -> str:
    return f"{count} image(s) or diagram(s) not extracted"


def equation_issue(count: int) -> str:
    return f"{count} equation(s) not extracted"


def parse_pptx(path: Path, max_pages: int) -> tuple[Page, ...]:
    try:
        presentation = Presentation(str(path))
    except Exception as e:  # untrusted input: any failure to parse becomes a user-facing error
        raise IngestError(f"Could not open {path.name}: {e}") from e

    slides = list(presentation.slides)
    if len(slides) > max_pages:
        raise TooManyPagesError(f"{path.name} has {len(slides)} slides; the limit is {max_pages}")
    slide_area = (presentation.slide_width or 0) * (presentation.slide_height or 0)
    # Hidden slides are kept, so numbers match what PowerPoint shows.
    return tuple(
        _parse_slide(slide, number=n, slide_area=slide_area)
        for n, slide in enumerate(slides, start=1)
    )


def _parse_slide(slide: Slide, *, number: int, slide_area: int) -> Page:
    shapes = _reading_order(_flatten(slide.shapes))
    blocks = [block for shape in shapes if (block := _to_block(shape))]
    if slide.has_notes_slide and (notes_frame := slide.notes_slide.notes_text_frame):
        notes = clean_text(notes_frame.text)
        if notes:
            blocks.append(Block(kind=BlockKind.NOTES, text=notes))

    issues: list[str] = []
    if slide.element.get("show") == "0":
        issues.append(HIDDEN_ISSUE)
    visuals = sum(1 for shape in shapes if _is_large_visual(shape, slide_area))
    if visuals:
        issues.append(visual_issue(visuals))
    equations = len(slide.element.findall(f".//{{{_MATH_NS}}}oMath"))
    if equations:
        issues.append(equation_issue(equations))

    return Page(number=number, blocks=tuple(blocks), issues=tuple(issues))


def _flatten(shapes: Iterable[BaseShape]) -> Iterator[BaseShape]:
    """Yield the shapes inside groups instead of the groups themselves."""
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _flatten(shape.shapes)
        else:
            yield shape


def _reading_order(shapes: Iterable[BaseShape]) -> list[BaseShape]:
    # Shapes are stored in stacking order, not reading order: sort title first, then top to
    # bottom and left to right. Shapes inside a group use the group's own coordinates, so
    # their order is approximate; good enough for a baseline.
    return sorted(shapes, key=lambda s: (not _is_title(s), s.top or 0, s.left or 0))


def _to_block(shape: BaseShape) -> Block | None:
    if isinstance(shape, GraphicFrame) and shape.has_table:
        rows = (
            " | ".join(" ".join(cell.text.split()) for cell in row.cells)
            for row in shape.table.rows
        )
        kind, text = BlockKind.TABLE, "\n".join(rows)
    elif isinstance(shape, Shape) and shape.has_text_frame:
        # Indent nested bullets, so the structure of the slide survives as plain text.
        lines = ("  " * p.level + p.text for p in shape.text_frame.paragraphs)
        kind = BlockKind.TITLE if _is_title(shape) else BlockKind.TEXT
        text = "\n".join(lines)
    else:
        return None

    text = clean_text(text)
    return Block(kind=kind, text=text) if text else None


def _is_title(shape: BaseShape) -> bool:
    return shape.is_placeholder and shape.placeholder_format.type in _TITLE_TYPES


def _is_large_visual(shape: BaseShape, slide_area: int) -> bool:
    """A picture, chart, SmartArt or media object big enough to be content, not decoration."""
    is_visual = isinstance(shape, Picture) or (
        isinstance(shape, GraphicFrame) and not shape.has_table
    )
    area = (shape.width or 0) * (shape.height or 0)
    return is_visual and slide_area > 0 and area / slide_area >= MIN_VISUAL_SHARE
